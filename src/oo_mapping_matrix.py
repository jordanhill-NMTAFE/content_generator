from itertools import chain
import os
from os import environ as env
import click
from docx import Document
from docx.shared import Pt, Inches
from docx.table import Table, _Cell
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.document import Document as _Document
from docx.section import _Header
from pathlib import Path
from typing import List, Dict
from src.utils.markdown import parse_md
from src.utils.uoc_api import UnitOfCompetency
from docx.enum.text import WD_ALIGN_PARAGRAPH
from frontmatter import Post

# Constants
ASSESSMENTS = Path("2 KAD/5 Assess Tool/")
MAPPING_MATRIX = Path("2 KAD/7 Assess Mapping Matrix/")
TEMPLATE = Path("templates/Assessment Mapping Matrix (F122A8).docx")
OUTPUT_FILE = Path("Assessment Mapping Matrix (F122A8).docx")
ROOT = Path(__file__).parent.parent.resolve()

# Environment Variables
COURSE_CONTENT = Path(env["COURSE_CONTENT"]).resolve()
OUTPUT_LOCATION = Path(env["OUTPUT_LOCATION"]).resolve()


class HeaderSection:
    def __init__(self, doc: _Document, mapping_matrix: Dict):
        self.doc = doc
        self.mapping_matrix = mapping_matrix

    def populate(self):
        unit = self.mapping_matrix.get("unit")
        header: _Header = self.doc.sections[0].header
        table_header: Table = header.tables[0]

        # Set qualification national codes and titles
        cell: _Cell = table_header.cell(0, 1)
        cell.text = f'{self.mapping_matrix.get("qualification")}'
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Set Unit national codes and titles
        cell: _Cell = table_header.cell(1, 1)
        cell.text = f'{unit.get("id")} {unit.get("name")}'
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


class ElementsSection:
    def __init__(self, doc: _Document, uoc: UnitOfCompetency):
        self.doc = doc
        self.uoc = uoc
        self.table = doc.tables[0]

    def populate(self):
        elements: dict = self.uoc.data.elements_and_criteria
        for element_index, (element, criteria) in enumerate(elements.items()):
            element_header = next(
                (
                    index
                    for index, cell in enumerate(self.table.column_cells(0))
                    if f"Element {element_index + 1}" in cell.text
                )
            )
            for index, criterium in enumerate(criteria.keys()):
                self.table.cell(element_header + 1 + index, 0).text = criterium


class KnowledgeEvidenceSection:
    def __init__(self, doc: _Document, uoc: UnitOfCompetency):
        self.doc = doc
        self.uoc = uoc
        self.table = doc.tables[0]

    def populate(self):
        knowledge_elements = self.uoc.data.knowledge_evidence
        knowledge_header = next(
            (
                index
                for index, cell in enumerate(self.table.column_cells(0))
                if "Required Knowledge or Knowledge Evidence" in cell.text
            ),
            None,
        )
        if knowledge_header is None:
            print("Knowledge Evidence header not found in the document.")
            return

        blurb, KE = next(iter(knowledge_elements.items()))
        for index, (element, subelements) in enumerate(KE.items()):
            cell = self.table.cell(knowledge_header + 1 + index, 0)
            cell.text = element
            cell.paragraphs[0].runs[0].bold = True
            if subelements:
                for sub_element in subelements:
                    paragraph = cell.add_paragraph()
                    paragraph.paragraph_format.left_indent = Inches(0.5)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.add_run(sub_element)


class PerformanceEvidenceSection:
    def __init__(self, doc: _Document, uoc: UnitOfCompetency):
        self.doc = doc
        self.uoc = uoc
        self.table = doc.tables[0]

    def populate(self):
        performance_elements = self.uoc.data.performance_evidence
        performance_header = next(
            (
                index
                for index, cell in enumerate(self.table.column_cells(0))
                if "Required Skills or Performance Evidence" in cell.text
            ),
            None,
        )
        if performance_header is None:
            print("Performance Evidence header not found in the document.")
            return

        counter = 0
        for index, (element, subelements) in enumerate(performance_elements.items()):
            cell = self.table.cell(performance_header + 1 + index + counter, 0)
            cell.text = element
            if ":" in cell.text:
                cell.paragraphs[0].runs[0].bold = True
            if subelements:
                for sub_element, indented_elements in subelements.items():
                    counter += 1
                    cell = self.table.cell(performance_header + 1 + index + counter, 0)
                    cell.text = sub_element
                    for indented_element in indented_elements:
                        paragraph = cell.add_paragraph()
                        paragraph.paragraph_format.left_indent = Inches(0.5)
                        paragraph.paragraph_format.space_after = Pt(0)
                        paragraph.add_run(indented_element)


class AssessmentConditionsSection:
    def __init__(self, doc: _Document, uoc: UnitOfCompetency):
        self.doc = doc
        self.uoc = uoc
        self.table = doc.tables[
            0
        ]  # Assumes the assessment conditions are in the first table

    def get_column_cells(self, column_idx: int) -> List[_Cell]:
        """Helper method to retrieve all cells in a specified column."""
        return [row.cells[column_idx] for row in self.table.rows]

    def populate(self):
        assessment_conditions = self.uoc.data.assessment_conditions
        column_index = 0  # Assuming the first column has the labels

        # Find the header row for "Assessment Conditions"
        ac_header = next(
            (
                index
                for index, cell in enumerate(self.get_column_cells(column_index))
                if "Assessment Conditions" in cell.text
            ),
            None,
        )
        if ac_header is None:
            print("Assessment Conditions header not found in the document.")
            return

        rows = []
        for element, subelements in assessment_conditions.items():
            row_index = ac_header + 1 + len(rows)
            rows.append(row_index)

            # Populate the first column with element text
            cell = self.table.cell(row_index, column_index)
            cell.text = element
            if ":" in element:
                for run in cell.paragraphs[0].runs:
                    run.bold = True

            # Add subelements as indented paragraphs
            if subelements:
                for sub_element in subelements:
                    paragraph = cell.add_paragraph()
                    paragraph.paragraph_format.left_indent = Inches(0.5)
                    paragraph.paragraph_format.space_after = Pt(0)
                    run = paragraph.add_run(sub_element)
                    run.bold = False  # Assuming subelements should not be bold

            # Merge cells across columns (from column 1 to the last column)
            target_cell = self.table.cell(row_index, 1)
            for col_idx in range(2, len(self.table.columns)):
                target_cell.merge(self.table.cell(row_index, col_idx))

        # Merge cells vertically in both columns 0 and 1
        if rows:
            rows_sorted = sorted(rows)
            first_row = rows_sorted.pop(0)

            for other_row in rows_sorted:
                # Merge column 1 cells vertically
                self.table.cell(first_row, 1).merge(self.table.cell(other_row, 1))
                # Merge column 0 cells vertically
                self.table.cell(first_row, 0).merge(self.table.cell(other_row, 0))

            # Update the merged cell text by replacing "must" phrases
            merged_cell = self.table.cell(first_row, 1)
            merged_text = (
                "\n".join(
                    (
                        item if isinstance(item, str) else "\n".join(item.keys())
                        for item in chain(*assessment_conditions.items())
                        if len(item) > 0
                    )
                )
                .replace("must be", "are")
                .replace("must", "always")
            )
            merged_cell.text = merged_text

            # Optionally, apply formatting to the merged text
            for paragraph in merged_cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.bold = False  # Adjust as needed


class AssessmentsSection:
    def __init__(
        self,
        doc: _Document,
        assessments: List[Post],
        uoc: UnitOfCompetency,
        unit_id: str,
    ):
        self.doc = doc
        self.assessments = assessments
        self.uoc = uoc
        self.unit_id = unit_id
        self.table = doc.tables[0]

    def populate(self):
        elements = self.uoc.data.elements_and_criteria
        knowledge_elements = self.uoc.data.knowledge_evidence

        for assessment_index, assessment in enumerate(self.assessments):
            cell = self.table.cell(0, assessment_index + 1)
            paragraph = cell.paragraphs[0]
            paragraph.clear()
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            # Set up columns
            paragraph.text = f"Assessment Task {assessment_index + 1}"
            paragraph.style = self.doc.styles["Heading 3"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # Set up Assessment Title
            cell = self.table.cell(1, assessment_index + 1)
            cell.text = assessment.get("name", "")
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # Mapping
            mapping = assessment.get("mapping", []) or []

            # Map Elements
            self.map_elements(elements, mapping, assessment_index)

            # Map Knowledge Evidence
            self.map_knowledge_evidence(knowledge_elements, mapping, assessment_index)

            # Map Performance Evidence
            self.map_performance_evidence(mapping, assessment_index)

    def map_elements(self, elements: dict, mapping: list, assessment_index: int):
        for element_index, (element, criteria) in enumerate(elements.items()):
            element_header = next(
                (
                    index
                    for index, cell in enumerate(self.table.column_cells(0))
                    if f"Element {element_index + 1}" in cell.text
                )
            )
            for index, criterium in enumerate(criteria.keys()):
                key: float = float(criterium[:3])
                question_mapping: str = ", ".join(
                    (
                        str(question_index + 1)
                        for question_index, question in enumerate(mapping)
                        if key
                        in (
                            ((question or {}).get("criteria") or {}).get(self.unit_id)
                            or []
                        )
                    )
                )
                cell = self.table.cell(element_header + 1 + index, 1 + assessment_index)
                cell.text = question_mapping

                # Center align the paragraph in the cell
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def map_knowledge_evidence(
        self, knowledge_elements: dict, mapping: list, assessment_index: int
    ):
        knowledge_header = next(
            (
                index
                for index, cell in enumerate(self.table.column_cells(0))
                if "Required Knowledge or Knowledge Evidence" in cell.text
            ),
            None,
        )
        if knowledge_header is None:
            print("Knowledge Evidence header not found in the document.")
            return

        blurb, KE = next(iter(knowledge_elements.items()))
        for knowledge_index, (element, subelements) in enumerate(KE.items()):
            cell = self.table.cell(
                knowledge_header + 1 + knowledge_index, 1 + assessment_index
            )
            key: int = knowledge_index + 1
            question_mapping: str = ", ".join(
                (
                    str(question_index + 1)
                    for question_index, question in enumerate(mapping)
                    if key
                    in (
                        ((question or {}).get("knowledge") or {}).get(self.unit_id)
                        or []
                    )
                )
            )
            cell.text = question_mapping

            # Center align the paragraph in the cell
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def map_performance_evidence(self, mapping: list, assessment_index: int):
        performance_header = next(
            (
                index
                for index, cell in enumerate(self.table.column_cells(0))
                if "including evidence of the ability to:" in cell.text
            ),
            None,
        )
        if performance_header is None:
            print("Performance Evidence header not found in the document.")
            return

        skills_header = next(
            (
                index
                for index, cell in enumerate(self.table.column_cells(0))
                if "In the course of the above, the candidate must:" in cell.text
            ),
            None,
        )
        if skills_header is None:
            print("Skills header not found in the document.")
            return

        # Mapping Performance Evidence
        for performance_index, element_number in enumerate(
            chain(
                *[
                    ((question or {}).get("performance") or {}).get(self.unit_id) or []
                    for question in mapping
                ]
            )
        ):
            cell = self.table.cell(
                performance_header + 1 + performance_index, 1 + assessment_index
            )
            key: int = performance_index + 1
            question_mapping: str = ", ".join(
                (
                    str(question_index + 1)
                    for question_index, question in enumerate(mapping)
                    if key
                    in (
                        ((question or {}).get("performance") or {}).get(self.unit_id)
                        or []
                    )
                )
            )
            cell.text = question_mapping

            # Center align the paragraph in the cell
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Mapping Skills
        for skills_index, element_number in enumerate(
            chain(
                *[
                    ((question or {}).get("skills") or {}).get(self.unit_id) or []
                    for question in mapping
                ]
            )
        ):
            cell = self.table.cell(
                skills_header + 1 + skills_index, 1 + assessment_index
            )
            key: int = skills_index + 1
            question_mapping: str = ", ".join(
                (
                    str(question_index + 1)
                    for question_index, question in enumerate(mapping)
                    if key
                    in (((question or {}).get("skills") or {}).get(self.unit_id) or [])
                )
            )
            cell.text = question_mapping

            # Center align the paragraph in the cell
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


class MappingMatrix:
    def __init__(self, course_directory: Path = None, output_location: Path = None):
        self.setup_environment(course_directory, output_location)
        self.assessments = self.course_directory / ASSESSMENTS
        self.unit_assessment_mapping = {}
        self.template = ROOT / TEMPLATE
        self.build_unit_assessment_mapping()

    def setup_environment(self, course_directory: Path, output_location: Path):
        os.environ["ROOT_DIR"] = str(Path(__file__).parent.parent.resolve())
        assert (
            "COURSE_CONTENT" in env or course_directory is not None
        ), "COURSE_CONTENT is undefined"
        assert (
            "OUTPUT_LOCATION" in env or output_location is not None
        ), "OUTPUT_LOCATION is undefined"
        self.course_directory = (
            course_directory or Path(env["COURSE_CONTENT"]).resolve()
        )
        self.output_location = output_location or Path(env["OUTPUT_LOCATION"]).resolve()

    def build_unit_assessment_mapping(self):
        for assessment in sorted(self.assessments.rglob("assessment.md")):
            if not assessment.is_file():
                continue

            markdown = parse_md(assessment)
            name = markdown.get("name")
            units = markdown.get("units")

            for unit in units:
                # Initialize unit mapping matrix if needed:
                self.unit_assessment_mapping.setdefault(
                    unit["id"],
                    {
                        "assessments": [],
                        "unit": unit,
                        "qualification": markdown.get(
                            "qualification_national_code_and_title"
                        ),
                    },
                )
                # Add assessment to unit mapping matrix
                self.unit_assessment_mapping.get(unit["id"]).get("assessments").append(
                    markdown
                )

    def process_unit(self, unit_id: str, mapping_matrix: Dict):
        doc: _Document = Document(self.template)

        # Unit of Competency object
        uoc: UnitOfCompetency = UnitOfCompetency(unit_id)

        # Populate Header Section
        header_section = HeaderSection(doc, mapping_matrix)
        header_section.populate()

        # Elements Section
        elements_section = ElementsSection(doc, uoc)
        elements_section.populate()

        # Knowledge Evidence Section
        knowledge_section = KnowledgeEvidenceSection(doc, uoc)
        knowledge_section.populate()

        # Performance Evidence Section
        performance_section = PerformanceEvidenceSection(doc, uoc)
        performance_section.populate()

        # Assessment Conditions Section
        assessment_conditions_section = AssessmentConditionsSection(doc, uoc)
        assessment_conditions_section.populate()

        # Assessments Section
        assessments: List[Post] = mapping_matrix.get("assessments")
        assessments_section = AssessmentsSection(doc, assessments, uoc, unit_id)
        assessments_section.populate()

        # Save the document
        output: Path = (
            self.output_location / MAPPING_MATRIX / (unit_id + " " + str(OUTPUT_FILE))
        )
        output.parent.mkdir(exist_ok=True, parents=True)
        doc.save(output)

    def generate(self):
        for unit_index, (unit_id, mapping_matrix) in enumerate(
            self.unit_assessment_mapping.items()
        ):
            self.process_unit(unit_id, mapping_matrix)


def mapping_matrix(course_directory: Path, output_location: Path):
    matrix = MappingMatrix(course_directory, output_location)
    matrix.generate()


@click.command()
def run_cli():
    """
    CLI tool to generate Assessment Mapping Matrix documents.
    """
    mapping_matrix(COURSE_CONTENT, OUTPUT_LOCATION)


if __name__ == "__main__":
    run_cli()
