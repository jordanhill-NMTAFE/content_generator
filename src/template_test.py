from docx import Document
from pathlib import Path


def print_template_summary(template_path: Path):
    doc = Document(template_path)

    print("==== Document Summary ====\n")

    # Iterate over sections and headers
    for section_index, section in enumerate(doc.sections):
        print(f"Section {section_index}:")
        header = section.header
        if header:
            print(f"  Header in Section {section_index}:")
            for para in header.paragraphs:
                print(f"    Paragraph: {para.text.strip()}")
            for table_idx, table in enumerate(header.tables):
                print(f"    Table {table_idx} in Header:")
                for row_idx, row in enumerate(table.rows):
                    print(f"      Row {row_idx}:")
                    for col_idx, cell in enumerate(row.cells):
                        cell_text = cell.text.strip().replace("\n", " ")
                        print(f"        Cell ({row_idx}, {col_idx}): {cell_text}")

    # Iterate over all tables in the document
    for table_idx, table in enumerate(doc.tables):
        print(f"\nTable {table_idx}:")
        for row_idx, row in enumerate(table.rows):
            print(f"  Row {row_idx}:")
            for col_idx, cell in enumerate(row.cells):
                cell_text = cell.text.strip().replace("\n", " ")
                print(f"    Cell ({row_idx}, {col_idx}): {cell_text}")


if __name__ == "__main__":
    # Replace with the actual path to your template file
    template_path = Path("templates/Assessment Mapping Matrix (F122A8).docx")
    if not template_path.is_file():
        print(f"Template file not found at {template_path}")
    else:
        print_template_summary(template_path)
