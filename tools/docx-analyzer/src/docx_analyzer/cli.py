"""CLI entry point and analysis logic for docx-analyzer."""

import argparse
import json
from datetime import datetime
from pathlib import Path


def _detect_table_edge_cases(doc, analysis):
    """
    Detect and report table structural edge cases that might cause issues.

    Edge cases include:
    - Rows with fewer cells than the table's column count
    - Merged cells that create accessibility issues
    - Inconsistent row structures
    - Content controls that affect cell structure
    """
    print(f"\n🔍 EDGE CASE DETECTION:")
    print("-" * 40)

    edge_cases_found = False

    for table_idx, table in enumerate(doc.tables):
        table_columns = len(table.columns)
        issues = []

        # Check each row for cell count mismatches
        for row_idx, row in enumerate(table.rows):
            actual_cells = len(row.cells)

            if actual_cells < table_columns:
                issues.append(
                    {
                        "type": "cell_count_mismatch",
                        "row": row_idx,
                        "expected": table_columns,
                        "actual": actual_cells,
                        "description": f"Row {row_idx} has {actual_cells} cells but table has {table_columns} columns",
                    }
                )

            # Check for accessibility issues
            try:
                # Try to access the last expected cell
                if table_columns > 1:
                    test_cell = table.cell(row_idx, table_columns - 1)
            except (IndexError, Exception) as e:
                issues.append(
                    {
                        "type": "cell_access_error",
                        "row": row_idx,
                        "column": table_columns - 1,
                        "error": str(e),
                        "description": f"Cannot access cell [{row_idx},{table_columns - 1}]: {e}",
                    }
                )

            # Check for content controls in cells that might affect structure
            for cell_idx, cell in enumerate(row.cells):
                try:
                    content_controls = cell._tc.xpath(
                        ".//w:sdt",
                        namespaces={
                            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
                        },
                    )
                    if content_controls:
                        issues.append(
                            {
                                "type": "content_control_detected",
                                "row": row_idx,
                                "column": cell_idx,
                                "count": len(content_controls),
                                "description": f"Cell [{row_idx},{cell_idx}] contains {len(content_controls)} content control(s)",
                            }
                        )
                except Exception:
                    # If content control detection fails, skip silently
                    pass

        if issues:
            edge_cases_found = True
            print(f"\n⚠️  TABLE {table_idx} STRUCTURAL ISSUES:")

            for issue in issues:
                if issue["type"] == "cell_count_mismatch":
                    print(
                        f"   • Row {issue['row']}: Expected {issue['expected']} cells, found {issue['actual']}"
                    )
                    print(
                        f"     → This may indicate merged cells or structural complexity"
                    )
                elif issue["type"] == "cell_access_error":
                    print(
                        f"   • Cannot access cell [{issue['row']},{issue['column']}]: {issue['error']}"
                    )
                    print(
                        f"     → Use table.cell(row, col) carefully or modify table structure"
                    )

            print(f"   💡 Consider:")
            print(f"      - Check if cells are merged and need to be split")
            print(f"      - Verify template structure matches expected layout")
            print(f"      - Use alternative cell access methods if needed")

            # Add edge case info to analysis data
            if "edge_cases" not in analysis:
                analysis["edge_cases"] = []
            analysis["edge_cases"].append({"table_index": table_idx, "issues": issues})

    if not edge_cases_found:
        print(
            "✅ No structural edge cases detected - all tables have consistent layouts"
        )

    print()


def analyze_word_document(args):
    """
    Comprehensive Word document analyzer utility.

    Analyzes structure, tables, paragraphs, and their relationships
    to help understand Word document templates and layouts.
    """
    from docx import Document

    doc_path = Path(args.document_path)

    if not doc_path.exists():
        print(f"Error: Document not found: {doc_path}")
        return 1

    if not doc_path.suffix.lower() == ".docx":
        print(f"Error: Expected .docx file, got: {doc_path.suffix}")
        return 1

    try:
        doc = Document(doc_path)
    except Exception as e:
        print(f"Error: Failed to load document: {e}")
        return 1

    # Create analysis data structure
    analysis = {
        "document_path": str(doc_path),
        "analysis_timestamp": datetime.now().isoformat(),
        "total_paragraphs": len(doc.paragraphs),
        "total_tables": len(doc.tables),
        "structure": [],
        "tables": [],
        "paragraphs": [],
    }

    print(f"=== WORD DOCUMENT ANALYSIS ===")
    print(f"Document: {doc_path}")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")
    print(f"Document elements: {len(doc.element.body)}")
    print("=" * 60)

    # Track table count for numbering
    table_count = 0

    # Analyze document structure (paragraphs and tables in order)
    if not args.tables_only:
        print(f"\n📋 DOCUMENT STRUCTURE (with context):")
        print("-" * 40)

        for i, element in enumerate(doc.element.body):
            if element.tag.endswith("p"):  # Paragraph
                para_text = ""
                for para in doc.paragraphs:
                    if para._element == element:
                        para_text = para.text.strip()
                        break

                if para_text:
                    print(f'📝 P{i}: "{para_text}"')
                    analysis["structure"].append(
                        {"type": "paragraph", "index": i, "content": para_text}
                    )
                    analysis["paragraphs"].append(
                        {"index": i, "content": para_text, "length": len(para_text)}
                    )
                elif not args.context_only:
                    print(f"📝 P{i}: [EMPTY]")
                    analysis["structure"].append(
                        {"type": "paragraph", "index": i, "content": ""}
                    )

            elif element.tag.endswith("tbl"):  # Table
                print(f"\n📊 TABLE {table_count} (follows P{i - 1})")

                table = doc.tables[table_count]
                rows = len(table.rows)
                cols = len(table.rows[0].cells) if table.rows else 0

                print(f"    📐 Dimensions: {rows} rows × {cols} columns")

                table_data = {
                    "index": table_count,
                    "follows_paragraph": i - 1,
                    "rows": rows,
                    "columns": cols,
                    "cells": [],
                }

                # Show table content unless context-only mode
                if not args.context_only:
                    for row_idx, row in enumerate(table.rows):
                        print(f"    📋 Row {row_idx}:")
                        row_data = []
                        for cell_idx, cell in enumerate(row.cells):
                            cell_text = cell.text.strip()
                            if cell_text:
                                display_text = cell_text[:60] + (
                                    "..." if len(cell_text) > 60 else ""
                                )
                                print(
                                    f'        [{row_idx},{cell_idx}]: "{display_text}"'
                                )
                                row_data.append(cell_text)
                            else:
                                print(f"        [{row_idx},{cell_idx}]: [EMPTY]")
                                row_data.append("")
                        table_data["cells"].append(row_data)
                else:
                    # Just show first row for context
                    if table.rows:
                        first_row_cells = []
                        for cell in table.rows[0].cells:
                            cell_text = cell.text.strip()[:40]
                            if cell_text:
                                first_row_cells.append(cell_text)
                        if first_row_cells:
                            print(f"    📋 Headers: {' | '.join(first_row_cells)}")

                analysis["structure"].append(
                    {
                        "type": "table",
                        "index": table_count,
                        "follows_paragraph": i - 1,
                        "dimensions": f"{rows}×{cols}",
                    }
                )
                analysis["tables"].append(table_data)

                table_count += 1
                print()

    # Edge case detection for table structural issues
    _detect_table_edge_cases(doc, analysis)

    # Table-only analysis
    if args.tables_only:
        print(f"\n📊 TABLES SUMMARY:")
        print("-" * 40)

        for i, table in enumerate(doc.tables):
            rows = len(table.rows)
            cols = len(table.rows[0].cells) if table.rows else 0
            print(f"\nTable {i}: {rows} rows × {cols} columns")

            table_data = {"index": i, "rows": rows, "columns": cols, "cells": []}

            for row_idx, row in enumerate(table.rows):
                print(f"  Row {row_idx}:")
                row_data = []
                for cell_idx, cell in enumerate(row.cells):
                    cell_text = cell.text.strip()
                    if cell_text:
                        display_text = cell_text[:80] + (
                            "..." if len(cell_text) > 80 else ""
                        )
                        print(f'    [{row_idx},{cell_idx}]: "{display_text}"')
                        row_data.append(cell_text)
                    else:
                        print(f"    [{row_idx},{cell_idx}]: [EMPTY]")
                        row_data.append("")
                table_data["cells"].append(row_data)
            analysis["tables"].append(table_data)

        # Edge case detection for tables-only mode
        _detect_table_edge_cases(doc, analysis)

    # Export analysis if requested
    if args.export:
        export_path = Path(args.export)
        try:
            if export_path.suffix.lower() == ".json":
                with open(export_path, "w", encoding="utf-8") as f:
                    json.dump(analysis, f, indent=2, ensure_ascii=False)
                print(f"\n💾 Analysis exported to: {export_path}")
            else:
                # Export as text
                with open(export_path, "w", encoding="utf-8") as f:
                    f.write(f"Word Document Analysis\n")
                    f.write(f"Document: {doc_path}\n")
                    f.write(f"Generated: {analysis['analysis_timestamp']}\n")
                    f.write(f"Paragraphs: {analysis['total_paragraphs']}\n")
                    f.write(f"Tables: {analysis['total_tables']}\n\n")

                    for item in analysis["structure"]:
                        if item["type"] == "paragraph":
                            f.write(f"P{item['index']}: {item['content']}\n")
                        else:
                            f.write(f"TABLE {item['index']}: {item['dimensions']}\n")

                    f.write(f"\n\nDetailed Table Data:\n")
                    for table in analysis["tables"]:
                        f.write(f"\nTable {table['index']}:\n")
                        for row_idx, row in enumerate(table["cells"]):
                            f.write(f"  Row {row_idx}: {row}\n")

                print(f"\n💾 Analysis exported to: {export_path}")
        except Exception as e:
            print(f"\n❌ Export failed: {e}")
            return 1

    print(f"\n✅ Analysis complete!")
    print(f"\n💡 TIP: Use this analysis to understand:")
    print(f"   • Which tables are for what purpose (check preceding paragraphs)")
    print(f"   • Table structure and available cells")
    print(f"   • Document flow and layout")
    print(f"   • Export with --export to save analysis for reference")

    return 0


def main():
    """Entry point for the docx CLI tool."""
    parser = argparse.ArgumentParser(
        prog="docx",
        description="Word document analysis tools",
    )
    subparsers = parser.add_subparsers(dest="command", help="Available commands")

    # analyze subcommand
    analyze_parser = subparsers.add_parser(
        "analyze",
        help="Analyze Word document structure (tables, paragraphs, context)",
    )

    analyze_parser.add_argument(
        "document_path",
        type=str,
        help="Path to the Word document (.docx) to analyze",
    )

    analyze_parser.add_argument(
        "--tables-only",
        "-t",
        action="store_true",
        help="Show only table structures (skip paragraphs)",
    )

    analyze_parser.add_argument(
        "--context-only",
        "-c",
        action="store_true",
        help="Show only paragraph context (skip detailed table content)",
    )

    analyze_parser.add_argument(
        "--export",
        "-e",
        type=str,
        help="Export analysis to file (txt or json format)",
    )

    args = parser.parse_args()

    if not args.command:
        parser.print_help()
        return 1

    if args.command == "analyze":
        return analyze_word_document(args)

    parser.print_help()
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
